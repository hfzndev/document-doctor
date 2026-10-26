"""Build the sample templates in fixtures/ used by tests and the demo:
sample-letter.docx (agreement) and sample-invoice.docx (numbered, dated letter).

    python3 scripts/make-sample.py
"""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


def add_field(paragraph, instr):
    run = paragraph.add_run()
    for kind, text in (("begin", None), (None, instr), ("separate", None), (None, "1"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
            run._r.append(el)
        elif text == instr:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = f" {instr} "
            run._r.append(el)
        else:
            t = OxmlElement("w:t")
            t.text = text
            run._r.append(t)
        run = paragraph.add_run()


def add_hyperlink(paragraph, url, text):
    part = paragraph.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                          is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1D4ED8")
    rpr.append(color)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


doc = Document()
section = doc.sections[0]

header = section.header.paragraphs[0]
header.text = "NORFLABS · Internal"
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.add_run("Page ")
add_field(footer, "PAGE")
footer.add_run(" of ")
add_field(footer, "NUMPAGES")

doc.add_paragraph("Service Agreement", style="Title")
doc.add_paragraph("Between Norflabs and {Client Name}", style="Subtitle")
doc.add_heading("1. Scope of work", level=1)
p = doc.add_paragraph("Norflabs will deliver ")
bold = p.add_run("the website redesign")
bold.bold = True
p.add_run(" described in the attached proposal.")
doc.add_heading("1.1 Deliverables", level=2)
doc.add_paragraph("Design files", style="List Bullet")
doc.add_paragraph("Production build", style="List Bullet")
doc.add_heading("2. Fees", level=1)
table = doc.add_table(rows=2, cols=2)
table.style = "Table Grid"
table.cell(0, 0).text = "Item"
table.cell(0, 1).text = "Amount"
table.cell(1, 0).text = "Total"
table.cell(1, 1).text = "Rp 10.000.000"
p = doc.add_paragraph("Questions? Visit ")
add_hyperlink(p, "https://norflabs.com", "norflabs.com")
p.add_run(".")
signature = doc.add_paragraph("Hafiz Norman")
signature.runs[0].font.size = Pt(12)
signature.runs[0].font.color.rgb = RGBColor(0x33, 0x33, 0x33)
doc.add_paragraph("")

fixtures = Path(__file__).resolve().parent.parent / "fixtures"
doc.save(fixtures / "sample-letter.docx")
print(fixtures / "sample-letter.docx")

# --- invoice: a document number, a date and a client that repeat across documents
inv = Document()
s = inv.sections[0]
s.header.paragraphs[0].text = "PT NORFLABS DIGITAL"
foot = s.footer.paragraphs[0]
foot.alignment = WD_ALIGN_PARAGRAPH.RIGHT
foot.add_run("Ref. 001/INV/X/2026 · page ")
add_field(foot, "PAGE")

inv.add_paragraph("INVOICE", style="Title")
p = inv.add_paragraph()
p.add_run("Nomor: ").bold = True
p.add_run("001/INV/X/2026")
p = inv.add_paragraph()
p.add_run("Jakarta, 26 Oktober 2026")
inv.add_paragraph("Kepada Yth. PT Contoh Sejahtera")
inv.add_heading("Rincian", level=1)
inv.add_paragraph("Pekerjaan desain ulang website sesuai penawaran.")
inv.add_paragraph("Total: Rp 10.000.000")
inv.save(fixtures / "sample-invoice.docx")
print(fixtures / "sample-invoice.docx")
